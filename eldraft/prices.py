"""Import the official credit prices from the game's own player export.

The EuroLeague Fantasy stats screen exports an .xlsx whose ``Quotation`` column is the
real credit value (CR) of every player - the same number the app charges you.  That is
the ground truth we would otherwise have to model, so when the sheet is present we bake
its prices straight into the board and mark them ``price_source="official"``.

The export is a stable, first-party format (``ID, Name, Surname, Team, Position, ...,
Quotation, ...``) so - unlike the fuzzy availability reader in :mod:`eldraft.excel_io` -
this parser matches its columns by name and keeps the join to the model's own players
strict: a wrong link would put another player's price on your budget line.
"""
from __future__ import annotations

import os
from typing import Dict, List, Optional

from . import ratings
from .config import norm_position
from .names import NameIndex, clean

# Header tokens we expect in the official export, matched after cleaning.
_COL = {
    "id":        ["id"],
    "name":      ["name"],           # first name
    "surname":   ["surname"],        # last name
    "team":      ["team"],
    "position":  ["position"],
    "quotation": ["quotation", "quot"],
}


def _find_columns(rows: List[list]) -> Dict[str, int]:
    """Locate the header row and map our fields to column indexes."""
    best: Dict[str, int] = {}
    for row in rows[:12]:
        mapping: Dict[str, int] = {}
        for j, cell in enumerate(row):
            h = clean(str(cell or ""))
            if not h:
                continue
            for field, toks in _COL.items():
                if field in mapping:
                    continue
                if any(h == t for t in toks):
                    mapping[field] = j
        # A valid header row has at least a surname and the quotation.
        if "surname" in mapping and "quotation" in mapping and len(mapping) > len(best):
            best = mapping
    return best


def read_quotations(path: str) -> List[dict]:
    """-> [{first, last, team, position, quotation, dunkest_id}] from the export."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".xlsx", ".xlsm", ".xltx"):
        raise ValueError("{}: expected the game's .xlsx export".format(path))
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    cols = _find_columns(rows)
    if "surname" not in cols or "quotation" not in cols:
        raise ValueError(
            "{}: could not find 'Surname' and 'Quotation' columns. Headers seen: {}".format(
                path, [str(c) for c in (rows[0] if rows else [])][:12]))

    # First data row is the one after the header we settled on.
    hdr_i = 0
    for i, row in enumerate(rows[:12]):
        if all(cols.get(f) is not None and cols[f] < len(row)
               and clean(str(row[cols[f]] or "")) in _COL[f]
               for f in ("surname", "quotation")):
            hdr_i = i
            break

    def q(v) -> Optional[float]:
        if v in (None, "", "-"):
            return None
        try:
            return float(str(v).replace(",", "."))
        except ValueError:
            return None

    out: List[dict] = []
    for row in rows[hdr_i + 1:]:
        def cell(f):
            j = cols.get(f)
            return row[j] if (j is not None and j < len(row)) else None

        last = cell("surname")
        quot = q(cell("quotation"))
        if not last or quot is None:
            continue
        out.append({
            "first": str(cell("name") or "").strip(),
            "last": str(last).strip(),
            "team": (str(cell("team")).strip() if cell("team") else None),
            "position": (str(cell("position")).strip() if cell("position") else None),
            "quotation": quot,
            "dunkest_id": (str(cell("id")).strip() if cell("id") else None),
        })
    return out


def apply_official_prices(board: dict, rows: List[dict], tag: str = "official") -> dict:
    """Overlay real quotations onto the board's players and coaches, matched by name.

    Ambiguous names (two players sharing a surname the sheet can't separate) are narrowed
    by position first, and dropped rather than guessed at if still tied.
    """
    players = board.get("players", [])
    coaches = board.get("coaches", [])
    everyone = players + coaches

    idx = NameIndex()
    for p in everyone:
        idx.add(p.get("first", ""), p.get("last", ""), p)

    report = {"rows": len(rows), "priced": 0, "coaches_priced": 0,
              "unmatched": [], "ambiguous": []}
    seen = set()
    for r in rows:
        hit = idx.lookup(r["first"], r["last"], max_tier=2)
        if hit is None:
            # Retry: narrow the loosest bucket by position before giving up.
            want = norm_position(r["position"]) if r.get("position") else None
            from .names import key_variants
            cands: List[dict] = []
            for k in key_variants(clean(r["first"]), clean(r["last"])):
                cands = idx.map.get(k, [])
                if cands:
                    break
            if want:
                narrowed = [c for c in cands if c.get("position") == want]
                if len(narrowed) == 1:
                    hit = narrowed[0]
            if hit is None:
                nm = (r["first"] + " " + r["last"]).strip()
                (report["ambiguous"] if len(cands) > 1 else report["unmatched"]).append(nm)
                continue
        if id(hit) in seen:
            continue
        seen.add(id(hit))
        hit["price"] = round(float(r["quotation"]), 1)
        hit["price_source"] = tag
        if hit.get("dunkest_id") is None and r.get("dunkest_id"):
            hit["dunkest_id"] = r["dunkest_id"]
        if hit.get("role") == "coach" or hit.get("position") == "H":
            report["coaches_priced"] += 1
        else:
            report["priced"] += 1

    # Refresh value ratings so fp_per_credit reflects the real prices.
    ratings.attach_value_ratings(players)
    for c in coaches:
        c["fp_per_credit"] = round(c["fp"] / c["price"], 3) if c.get("price") else None
    report["board_players"] = len(players)
    report["board_coaches"] = len(coaches)
    report["unpriced_players"] = [
        p["name"] for p in players if p.get("price_source") != tag]
    return report


def import_prices(board: dict, sheet_path: str, tag: str = "official") -> dict:
    """Read the export and apply it; returns the match report."""
    return apply_official_prices(board, read_quotations(sheet_path), tag=tag)
